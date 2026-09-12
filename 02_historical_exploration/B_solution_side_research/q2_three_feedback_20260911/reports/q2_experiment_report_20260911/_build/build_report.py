from pathlib import Path
import collections
import copy
import json
import math
import re
import statistics
import sys

HERE = Path(__file__).resolve().parent
REPORT = HERE.parent
ROOT = REPORT.parent.parent
sys.path.insert(0, str(HERE / 'vendor'))
from latex2mathml.converter import convert
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def read(p):
    return json.loads(p.read_text(encoding='utf-8'))


verified = read(ROOT / 'results/final_comparison/verified_comparison.json')
summary = read(ROOT / 'results/mainline_v1/summary.json')
rows = read(ROOT / 'results/mainline_v1/feedback_rows.json')
worlds = collections.defaultdict(dict)
for r in rows:
    worlds[r['world_id']][r['mode']] = r
pairs = [(v['old_precision'], v['open_search']) for v in worlds.values()]
assert len(pairs) == 480 and len(rows) == 1920
assert all(r['truth_covered'] for r in rows)
assert all(x['modes']['open_search']['upper_m'] < x['modes']['old_precision']['lower_m'] for x in verified)
assert sum(c['unresolved_boxes'] for x in verified for m in x['modes'].values() for c in m['cover_checks']) == 0


def interval(r):
    return f"[{math.floor(r['lower_m']*100)/100:.2f}, {math.ceil(r['upper_m']*100)/100:.2f}]"


def table(headers, data):
    return '\n'.join(['| ' + ' | '.join(headers) + ' |', '| ' + ' | '.join(['---']*len(headers)) + ' |'] + ['| ' + ' | '.join(str(x) for x in row) + ' |' for row in data])


verified_table = table(['上下文', '原精度点最坏半径区间 米', '扩展搜索点最坏半径区间 米'], [
    [r['context_id'].replace('r1-p3-survey-', '').replace('-ch', ' ch'), interval(r['modes']['old_precision']), interval(r['modes']['open_search'])]
    for r in verified
])
synthetic_table = table(['策略', '平均半径 米', '样本最大半径 米', '无信号记录数'], [
    [m, f"{summary['synthetic_secondary'][m]['mean_radius_upper_m']:.2f}",
     f"{summary['synthetic_secondary'][m]['sample_max_radius_upper_m']:.2f}", summary['synthetic_secondary'][m]['no_signal_count']]
    for m in ['old_fixed', 'old_precision', 'open_seed', 'open_search']
])
paired = []
for label, kind in [('新方案仍有方向', 'direction'), ('新方案无信号', 'no_signal')]:
    group = [(a,b) for a,b in pairs if b['feedback']['kind'] == kind]
    old = statistics.mean(a['radius_upper_m'] for a,b in group)
    new = statistics.mean(b['radius_upper_m'] for a,b in group)
    paired.append({'group': label, 'n': len(group), 'old_mean': old, 'new_mean': new,
                   'overall_contribution': sum(b['radius_upper_m']-a['radius_upper_m'] for a,b in group)/480})
paired_table = table(['按新反馈分组', '记录数', '旧均值 米', '新均值 米'], [[r['group'],r['n'], f"{r['old_mean']:.2f}", f"{r['new_mean']:.2f}"] for r in paired])

md_path = REPORT / '第二问实验报告_最终稿.md'
md = md_path.read_text(encoding='utf-8')
for placeholder, value in [('VERIFIED_TABLE', verified_table), ('SYNTHETIC_TABLE', synthetic_table), ('PAIRED_TABLE', paired_table)]:
    md = md.replace(f'<!-- {placeholder} -->', value)
md = md.replace('<!-- REPORT_END -->', '')
md_path.write_text(md, encoding='utf-8')
(HERE/'report_data_checks.json').write_text(json.dumps({'verified_contexts':len(verified), 'worlds':len(pairs), 'rows':len(rows), 'paired_decomposition':paired, 'cover_unresolved_boxes':0}, ensure_ascii=False, indent=2), encoding='utf-8')

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(.72), Inches(.72)
sec.left_margin, sec.right_margin = Inches(.78), Inches(.78)
sec.header_distance, sec.footer_distance = Inches(.3), Inches(.32)


def font_for(style, name, size, east='宋体', bold=False):
    style.font.name, style.font.size = name, Pt(size)
    style.font.color.rgb = RGBColor(0,0,0)
    style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts'); rpr.insert(0,fonts)
    fonts.set(qn('w:eastAsia'), east)


font_for(doc.styles['Normal'], 'Times New Roman', 11)
normal = doc.styles['Normal'].paragraph_format
normal.space_after, normal.line_spacing = Pt(5), 1.18
normal.widow_control = True
for style, size in [('Title',20), ('Heading 1',15), ('Heading 2',12)]:
    font_for(doc.styles[style], 'Arial', size, '微软雅黑', style!='Title')
    pf = doc.styles[style].paragraph_format
    pf.space_before, pf.space_after = Pt(13 if style!='Title' else 0), Pt(7)
    pf.keep_with_next = True
font_for(doc.styles['Subtitle'], 'Arial', 11, '微软雅黑')
doc.styles['Subtitle'].font.italic=False
font_for(doc.styles['Caption'], 'Times New Roman', 9.5)
for border in list(doc.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)

header = sec.header.paragraphs[0]
header.text = '第二问  三反馈测点选择与定位实验报告'
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for r in header.runs:
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(0,0,0)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('第 ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
footer.add_run(' 页')
for r in footer.runs: r.font.size=Pt(9)

transform = etree.XSLT(etree.parse(r'C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL'))
math_ns = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
equations = []


def add_equation(latex):
    mathml = convert(latex)
    omml = transform(etree.fromstring(mathml.encode())).getroot()
    if doc.paragraphs:
        doc.paragraphs[-1].paragraph_format.keep_with_next = True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.keep_together = True
    for mr in omml.findall('.//{%s}r' % math_ns):
        mt = mr.find('{%s}t' % math_ns)
        if mt is not None and (mt.text or '') in {'min','max','sup','inf','rad','near'}:
            mpr=mr.find('{%s}rPr' % math_ns)
            if mpr is None:
                mpr=etree.Element('{%s}rPr' % math_ns); mr.insert(0,mpr)
            sty=etree.SubElement(mpr,'{%s}sty' % math_ns); sty.set('{%s}val' % math_ns,'p')
            mt.text='\u2009'+mt.text+'\u2009'
        wrpr = OxmlElement('w:rPr')
        fonts = OxmlElement('w:rFonts'); fonts.set(qn('w:ascii'),'Cambria Math'); fonts.set(qn('w:hAnsi'),'Cambria Math')
        wrpr.append(fonts)
        size = OxmlElement('w:sz'); size.set(qn('w:val'),'22'); wrpr.append(size)
        mr.insert(0,wrpr)
    p._p.append(copy.deepcopy(omml))
    equations.append(latex)


def inline(p, text):
    for part in re.split(r'(\*\*.*?\*\*)', text):
        r=p.add_run(part[2:-2] if part.startswith('**') else part)
        r.bold=part.startswith('**')


def add_table(lines):
    matrix = [[s.strip() for s in row.strip().strip('|').split('|')] for row in lines]
    matrix.pop(1)
    cols = len(matrix[0])
    t=doc.add_table(rows=1,cols=cols)
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    if cols==4:
        ratios = [.29,.18,.22,.31] if '符号' in matrix[0][0] else [.28,.23,.25,.24]
    elif cols==3: ratios=[.28,.36,.36]
    else: ratios=[.32,.68]
    available=6.94
    for col,ratio in zip(t.columns,ratios): col.width=Inches(available*ratio)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
    t._tbl.tblPr.append(borders)
    for rowid, values in enumerate(matrix):
        row=t.rows[0] if rowid==0 else t.add_row()
        trpr=row._tr.get_or_add_trPr()
        no_split=OxmlElement('w:cantSplit'); trpr.append(no_split)
        if rowid==0:
            repeat=OxmlElement('w:tblHeader'); trpr.append(repeat)
        for colid,(c,text) in enumerate(zip(row.cells,values)):
            c.width=Inches(available*ratios[colid])
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=c._tc.get_or_add_tcPr()
            mar=OxmlElement('w:tcMar')
            for k,v in [('top','75'),('bottom','75'),('left','90'),('right','90')]:
                e=OxmlElement('w:'+k); e.set(qn('w:w'),v); e.set(qn('w:type'),'dxa'); mar.append(e)
            tcpr.append(mar)
            if rowid==0:
                sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E5EBEF'); tcpr.append(sh)
            p=c.paragraphs[0]
            p.paragraph_format.line_spacing=1.1
            p.paragraph_format.space_after=Pt(0)
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT if colid==0 or cols==2 else WD_ALIGN_PARAGRAPH.CENTER
            if rowid==0: p.paragraph_format.keep_with_next=True
            r=p.add_run(text); r.font.size=Pt(9.5); r.bold=rowid==0
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(0); p.paragraph_format.space_before=Pt(0); p.paragraph_format.line_spacing=1
    p.add_run().font.size=Pt(3)


lines=md.splitlines(); i=0
while i<len(lines):
    line=lines[i].strip()
    if not line:
        i+=1; continue
    if line=='$$':
        i+=1; eq=[]
        while i<len(lines) and lines[i].strip()!='$$': eq.append(lines[i]); i+=1
        add_equation('\n'.join(eq)); i+=1; continue
    if line.startswith('```'):
        i+=1
        while i<len(lines) and not lines[i].startswith('```'):
            p=doc.add_paragraph()
            p.paragraph_format.space_after=Pt(2)
            p.paragraph_format.line_spacing=1.05
            r=p.add_run(lines[i]); r.font.name='Consolas'; r.font.size=Pt(8.5)
            i+=1
        i+=1; continue
    if line.startswith('|'):
        group=[]
        while i<len(lines) and lines[i].strip().startswith('|'): group.append(lines[i]); i+=1
        add_table(group); continue
    if line.startswith('# '):
        p=doc.add_paragraph(line[2:],'Title')
    elif line.startswith('## '):
        p=doc.add_paragraph(line[3:],'Heading 1')
    elif line.startswith('### '):
        p=doc.add_paragraph(re.sub(r'^(\d+)\.(\d+) ',r'\1 \2 ',line[4:]),'Heading 2')
    elif line=='原理分析 模型建立 模型求解 结果检验':
        p=doc.add_paragraph(line,'Subtitle')
    else:
        p=doc.add_paragraph(); inline(p,line)
        if re.match(r'^\d+\. ',line):
            p.paragraph_format.left_indent=Inches(.16)
            p.paragraph_format.first_line_indent=Inches(-.16)
        if line.startswith('[') and re.match(r'^\[\d+\]',line):
            for r in p.runs: r.font.size=Pt(9)
    i+=1

doc.core_properties.title='第二问三反馈测点选择与定位实验报告'
doc.core_properties.subject='三反馈测点选择模型和最新实验结果'
doc.core_properties.author=''
doc.core_properties.keywords='第二问 最小包围圆 三反馈 极小极大'
for border in list(doc.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
out=REPORT/'第二问实验报告_最终稿.docx'
doc.save(out)
(HERE/'equations.json').write_text(json.dumps(equations,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'docx':str(out),'markdown':str(md_path),'equations':len(equations),'tables':len(doc.tables),'paragraphs':len(doc.paragraphs)},ensure_ascii=False))
