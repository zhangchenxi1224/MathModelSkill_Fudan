from pathlib import Path
import sys,re,copy,json
WORK=Path(__file__).resolve().parent
sys.path.insert(0,str(WORK/'python_deps'))
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from latex2mathml.converter import convert

ROOT=Path(__file__).resolve().parents[1]
MD=ROOT/'P3_P4_模型与算法实验报告.md'
DOC=ROOT/'P3_P4_模型与算法实验报告.docx'
XSL=etree.XSLT(etree.parse('C:/Program Files/Microsoft Office/root/Office16/MML2OMML.XSL'))
doc=Document();sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(1.85);sec.right_margin=Cm(1.85)
sec.header_distance=Cm(.7);sec.footer_distance=Cm(.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
    st=doc.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体' if name in ['Normal','Caption'] else '黑体')
normal=doc.styles['Normal'];normal.font.size=Pt(11);normal.paragraph_format.line_spacing=1.3
normal.paragraph_format.space_after=Pt(5);normal.paragraph_format.widow_control=True
doc.styles['Title'].font.size=Pt(22);doc.styles['Title'].paragraph_format.space_after=Pt(12)
for name,size,before,after in [('Heading 1',16,13,8),('Heading 2',13,10,6),('Heading 3',11.5,8,5)]:
    st=doc.styles[name];st.font.size=Pt(size);st.font.bold=True;st.paragraph_format.space_before=Pt(before);st.paragraph_format.space_after=Pt(after);st.paragraph_format.keep_with_next=True
doc.styles['Caption'].font.size=Pt(9.5);doc.styles['Caption'].font.italic=False
doc.styles['Caption'].paragraph_format.space_after=Pt(8)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run();f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');r._r.addnext(f)
doc.core_properties.title='第三问与第四问全反馈搜索清除模型及实验报告'
doc.core_properties.subject='P3局部两步前瞻与P4三反馈期望成本决策'
doc.core_properties.author=''
equations=0
def mathxml(tex,display=False):
    global equations
    equations+=1
    result=XSL(etree.fromstring(convert(tex,display='block' if display else 'inline').encode('utf-8')))
    node=result.getroot()
    for r in node.iter(qn('m:r')):
        wr=OxmlElement('w:rPr');font=OxmlElement('w:rFonts');font.set(qn('w:ascii'),'Cambria Math');font.set(qn('w:hAnsi'),'Cambria Math');wr.append(font)
        sz=OxmlElement('w:sz');sz.set(qn('w:val'),'21' if display else '22');wr.append(sz);r.insert(0,wr)
    return node

def inline(p,text):
    # Native editable Office equations, ordinary text and code names.
    for tok in re.split(r'(\$[^$]+\$|`[^`]+`|\*\*[^*]+\*\*)',text):
        if not tok:continue
        if tok.startswith('$') and tok.endswith('$'):
            node=mathxml(tok[1:-1]);
            if node.tag==qn('m:oMathPara'):
                for child in list(node):
                    if child.tag==qn('m:oMath'):p._p.append(child)
            else:p._p.append(node)
        else:
            code=tok.startswith('`') and tok.endswith('`');bold=tok.startswith('**') and tok.endswith('**')
            r=p.add_run(tok[1:-1] if code else tok[2:-2] if bold else tok)
            if code:r.font.name='Consolas';r.font.size=Pt(9)
            if bold:r.bold=True

def table(lines):
    rows=[[v.strip() for v in re.split(r'(?<!\\)\|',ln.strip().strip('|'))] for ln in lines]
    rows=[r for r in rows if not all(re.fullmatch(r':?-+:?',v) for v in r)]
    n=len(rows[0]);t=doc.add_table(rows=1,cols=n);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    total=17.3
    header=rows[0]
    if n==2:widths=[8.5,8.8] if header[0]=='路径' else [4.1,13.2]
    elif n==3:widths=[1.8,7.2,8.3]
    elif '官方案例编码' in header:widths=[5.25,1.9,1.6,3.0,3.0,3.1]
    elif '全清' in header:widths=[3.4,2.0,3.0,3.0,2.1,3.8]
    elif n==6:widths=[3.4,1.55,3.0,3.0,2.1,4.25]
    elif n==5:widths=[4.4,3.225,3.225,3.225,3.225]
    else:widths=[total/n]*n
    scale=total/sum(widths);widths=[w*scale for w in widths]
    for i,w in enumerate(widths):t.columns[i].width=Cm(w)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for tag in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+tag);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    pr.append(borders)
    for ri,row in enumerate(rows):
        cells=t.rows[0].cells if ri==0 else t.add_row().cells
        for j,value in enumerate(row):
            cell=cells[j];cell.width=Cm(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=cell._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
            for side,num in [('top',85),('bottom',85),('left',95),('right',95)]:
                a=OxmlElement('w:'+side);a.set(qn('w:w'),str(num));a.set(qn('w:type'),'dxa');m.append(a)
            cp.append(m)
            if ri==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EEF2');cp.append(shade)
            p=cell.paragraphs[0];p.paragraph_format.line_spacing=1.12;p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.keep_with_next=ri<len(rows)-1
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 or n<=3 else WD_ALIGN_PARAGRAPH.CENTER
            inline(p,value)
            for r in p.runs:r.font.size=Pt(9.2);r.bold=ri==0
        trPr=t.rows[ri]._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trPr.append(cant)
        if ri==0:
            rep=OxmlElement('w:tblHeader');trPr.append(rep)
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1;p.add_run().font.size=Pt(2)

lines=MD.read_text(encoding='utf-8').splitlines();i=0
while i<len(lines):
    line=lines[i].strip()
    if not line:i+=1;continue
    if line.startswith('|'):
        block=[]
        while i<len(lines) and lines[i].strip().startswith('|'):block.append(lines[i]);i+=1
        table(block);continue
    if line=='$$':
        block=[];i+=1
        while i<len(lines) and lines[i].strip()!='$$':block.append(lines[i]);i+=1
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(8);p._p.append(mathxml('\n'.join(block),True));i+=1;continue
    if line.startswith('```'):
        i+=1
        while i<len(lines) and not lines[i].startswith('```'):
            p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.05
            r=p.add_run(lines[i]);r.font.name='Consolas';r.font.size=Pt(9.3);r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'等线')
            i+=1
        i+=1;doc.add_paragraph().paragraph_format.space_after=Pt(1);continue
    if line.startswith('!['):
        m=re.fullmatch(r'!\[(.*?)\]\((.*?)\)',line)
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
        p.add_run().add_picture(str(ROOT/m.group(2)),width=Cm(16.5))
        cap=doc.add_paragraph(m.group(1),style='Caption');cap.alignment=WD_ALIGN_PARAGRAPH.CENTER;i+=1;continue
    if line.startswith('#'):
        m=re.match(r'(#+)\s+(.+)',line);level=len(m.group(1));text=m.group(2)
        p=doc.add_paragraph(style='Title' if level==1 else f'Heading {min(level-1,3)}')
        if level==1:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        if text.startswith('2 第四问'):p.paragraph_format.page_break_before=True
        inline(p,text);i+=1;continue
    p=doc.add_paragraph();inline(p,line)
    if line.endswith('：'):p.paragraph_format.keep_with_next=True
    if re.fullmatch(r'\d{4}年\d+月\d+日',line):p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    i+=1
doc.save(DOC)
(WORK/'docx_build.json').write_text(json.dumps({'docx':str(DOC),'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'math_expressions':equations},ensure_ascii=False),encoding='utf-8')
print(DOC, 'equations',equations,'tables',len(doc.tables))
